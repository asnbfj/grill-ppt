#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
pptx_template - 用户自定义 PPT 母版加载器（Python / python-pptx 版）

与 scripts/template-loader.js 的定位一致：把 templates/ 下的用户 .pptx 母版
接进生成流程。区别在于本模块走的是**真正的母版继承**：

    prs = Presentation(template_path)   # 继承母版的 slideMaster / slideLayout / theme
    strip_slides(prs)                   # 再清空母版自带的幻灯片，避免污染输出

因此生成的 .pptx 在文件层面携带模板的母版、版式与主题（字体也随之继承），
而不是像 JS 版那样只能「提取主题配色」。同时本模块会把主题配色映射成
各生成脚本使用的调色板，并提供对比度保护。

零额外依赖：只用 python-pptx + 标准库 zipfile / xml.etree。
"""

from __future__ import print_function

import json
import os
import sys
import zipfile
import xml.etree.ElementTree as ET

from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor

try:  # 仅用于清空母版自带幻灯片
    from pptx.oxml.ns import qn
except ImportError:  # pragma: no cover
    qn = None

# ------------------------------------------------------------------ #
# 路径与常量
# ------------------------------------------------------------------ #

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATES_DIR = os.path.join(PROJECT_ROOT, 'templates')
REGISTRY_FILE = os.path.join(TEMPLATES_DIR, 'registry.json')

DEFAULT_FONT = 'Microsoft YaHei'
DEFAULT_SIZE = (13.333, 7.5)  # 16:9

A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'
THEME_NS = {'a': A_NS}


# ------------------------------------------------------------------ #
# 目录与登记表
# ------------------------------------------------------------------ #

def ensure_templates_dir():
    if not os.path.isdir(TEMPLATES_DIR):
        os.makedirs(TEMPLATES_DIR)
    return TEMPLATES_DIR


def load_registry():
    """读取 registry.json；缺失或非法 JSON 都退化为「无元数据」。"""
    try:
        with open(REGISTRY_FILE, 'r', encoding='utf-8') as fh:
            data = json.load(fh)
        return data if isinstance(data, dict) else {}
    except (IOError, OSError, ValueError):
        return {}


def list_templates():
    """列出 templates/ 下的所有 .pptx 母版，并合并 registry.json 元数据。"""
    ensure_templates_dir()
    registry = load_registry()

    result = []
    for name in sorted(os.listdir(TEMPLATES_DIR)):
        lower = name.lower()
        if not lower.endswith('.pptx') or name.startswith('~$'):
            continue
        meta = registry.get(name) or {}
        stem = os.path.splitext(name)[0]
        result.append({
            'file': name,
            'path': os.path.join(TEMPLATES_DIR, name),
            'name': meta.get('name') or stem,
            'style': meta.get('style') or '',
            'scene': meta.get('scene') or '',
            'desc': meta.get('desc') or '',
        })
    return result


def resolve_template(ref):
    """把模板引用（文件名 / 登记名 / 任意路径）解析成磁盘绝对路径。"""
    if not ref:
        raise ValueError('未提供模板名称或路径（--template 的值为空）')

    expanded = os.path.expanduser(ref)
    candidates = [
        expanded,
        os.path.abspath(expanded),
        os.path.join(TEMPLATES_DIR, expanded),
    ]
    for candidate in candidates:
        if os.path.isfile(candidate) and candidate.lower().endswith('.pptx'):
            return candidate

    for tpl in list_templates():
        stem = os.path.splitext(tpl['file'])[0]
        if ref in (tpl['file'], tpl['name'], stem):
            return tpl['path']

    available = '、'.join('{} ({})'.format(t['name'], t['file']) for t in list_templates())
    if not available:
        available = '（templates/ 目录为空，请先放入 .pptx 模板）'
    raise ValueError('找不到模板「{}」。可用模板：{}'.format(ref, available))


def print_templates(stream=None):
    stream = stream or sys.stdout
    templates = list_templates()
    if not templates:
        print('📁 templates/ 目录下暂无模板。把 .pptx 放进去即可，参见 templates/README.md。',
              file=stream)
        return
    print('📁 可用模板：', file=stream)
    for i, tpl in enumerate(templates, 1):
        detail = ' — '.join(x for x in (tpl['style'], tpl['scene']) if x)
        suffix = ' — ' + detail if detail else ''
        print('{}. {} ({}){}'.format(i, tpl['name'], tpl['file'], suffix), file=stream)


# ------------------------------------------------------------------ #
# 主题解析（zipfile + ElementTree，不依赖 python-pptx 内部结构）
# ------------------------------------------------------------------ #

def read_theme(template_path):
    """从 .pptx 里读 ppt/theme/theme*.xml，返回 {name, colors, fonts}。"""
    with zipfile.ZipFile(template_path) as zf:
        theme_names = [n for n in zf.namelist() if n.startswith('ppt/theme/theme') and n.endswith('.xml')]
        if not theme_names:
            raise ValueError('模板「{}」中未找到主题文件（ppt/theme/theme*.xml）'.format(template_path))
        xml_bytes = zf.read(sorted(theme_names)[0])

    root = ET.fromstring(xml_bytes)

    colors = {}
    clr_scheme = root.find('.//a:clrScheme', THEME_NS)
    if clr_scheme is not None:
        for child in clr_scheme:
            slot = child.tag.split('}')[-1]
            value = None
            srgb = child.find('a:srgbClr', THEME_NS)
            sysclr = child.find('a:sysClr', THEME_NS)
            if srgb is not None:
                value = srgb.get('val')
            elif sysclr is not None:
                value = sysclr.get('lastClr') or sysclr.get('val')
            if value:
                colors[slot] = value.upper()

    fonts = {'major': None, 'minor': None}
    for key, tag in (('major', 'majorFont'), ('minor', 'minorFont')):
        node = root.find('.//a:{}'.format(tag), THEME_NS)
        if node is not None:
            latin = node.find('a:latin', THEME_NS)
            if latin is not None and latin.get('typeface'):
                fonts[key] = latin.get('typeface')

    name = root.get('name') or ''
    return {'name': name, 'colors': colors, 'fonts': fonts}


# ------------------------------------------------------------------ #
# 颜色工具：对比度保护（呼应 SKILL.md「不让文字和背景对比度不足」）
# ------------------------------------------------------------------ #

def hex_to_rgb(value):
    value = value.lstrip('#').upper()
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def _luminance(value):
    channels = []
    for i in (0, 2, 4):
        v = int(value[i:i + 2], 16) / 255.0
        channels.append(v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4)
    return 0.2126 * channels[0] + 0.7152 * channels[1] + 0.0722 * channels[2]


def contrast_ratio(fg, bg):
    la, lb = _luminance(fg), _luminance(bg)
    hi, lo = max(la, lb), min(la, lb)
    return (hi + 0.05) / (lo + 0.05)


def ensure_contrast(color, bg, fallbacks=(), minimum=3.0):
    """保证 color 在 bg 上可读；不达标时依次尝试 fallbacks，最后按底色明暗兜底。"""
    def usable(c):
        return bool(c) and len(c.lstrip('#')) == 6

    if usable(color) and usable(bg) and contrast_ratio(color.lstrip('#'), bg.lstrip('#')) >= minimum:
        return color.lstrip('#').upper()
    for fb in fallbacks:
        if usable(fb) and usable(bg) and contrast_ratio(fb.lstrip('#'), bg.lstrip('#')) >= minimum:
            return fb.lstrip('#').upper()
    return '1F2937' if _luminance(bg.lstrip('#')) > 0.5 else 'F0F6FC'


def mix(a, b, t):
    """按比例混合两个 #RRGGBB 颜色。"""
    a, b = a.lstrip('#').upper(), b.lstrip('#').upper()
    out = ''
    for i in (0, 2, 4):
        av, bv = int(a[i:i + 2], 16), int(b[i:i + 2], 16)
        out += '{:02X}'.format(int(round(av + (bv - av) * t)))
    return out


# ------------------------------------------------------------------ #
# 主题配色 -> 语义角色
# ------------------------------------------------------------------ #

def build_roles(theme):
    """把主题色槽映射成语义角色（全部为 #RRGGBB 字符串）。无配色信息时返回 None。"""
    c = (theme or {}).get('colors') or {}
    if not c:
        return None

    def pick(*keys):
        for k in keys:
            if c.get(k):
                return c[k]
        return None

    bg_light = pick('lt1') or 'FFFFFF'
    bg_dark = pick('dk2', 'dk1') or '0D1117'

    primary = ensure_contrast(
        pick('accent1', 'dk2', 'dk1') or '1E2761', bg_light,
        [pick('accent1'), pick('accent2'), pick('accent3'), pick('dk2'), pick('dk1')], 3.0)

    text = ensure_contrast(
        pick('dk1', 'dk2') or '333333', bg_light,
        [pick('dk2'), pick('dk1'), pick('accent1')], 4.5)

    roles = {
        'primary': primary,
        'secondary': ensure_contrast(
            pick('accent2', 'lt2', 'accent1') or 'CADCFC', primary,
            [pick('accent2'), pick('lt2'), pick('lt1'), pick('accent3'), pick('accent4')], 3.0),
        'accent': ensure_contrast(
            pick('accent3', 'accent4', 'accent1', 'lt1') or 'FFFFFF', primary,
            [pick('accent3'), pick('accent4'), pick('lt1'), pick('accent2')], 3.0),
        'on_primary': ensure_contrast(
            'FFFFFF', primary, [pick('accent2'), pick('accent3'), pick('lt1')], 4.5),
        'bg_light': bg_light,
        'bg_dark': bg_dark,
        'text': text,
        'text_light': ensure_contrast(
            pick('dk2', 'dk1') or '6B7280', bg_light, ['6B7280'], 3.0),
        'on_dark': ensure_contrast(
            'FFFFFF', bg_dark, [pick('lt1'), pick('accent2')], 4.5),
        'slots': dict(c),
    }
    return roles


def make_palette(theme):
    """对外统一入口：解析主题 -> 语义角色（供各脚本映射调色板）。"""
    return build_roles(theme)


# ------------------------------------------------------------------ #
# 各脚本的调色板映射
# ------------------------------------------------------------------ #

def palette_generic(roles):
    """generate_full.py 的 COLORS 键位。"""
    return {
        'primary': hex_to_rgb(roles['primary']),
        'secondary': hex_to_rgb(roles['secondary']),
        'accent': hex_to_rgb(roles['accent']),
        'white': hex_to_rgb(roles['bg_light']),
        'text': hex_to_rgb(roles['text']),
        'gray': hex_to_rgb(roles['text_light']),
    }


def palette_grand(roles):
    """generate_grand.py 的深色大气键位。"""
    bg_dark = roles['bg_dark']
    return {
        'bg_dark': hex_to_rgb(bg_dark),
        'bg_gradient': hex_to_rgb(mix(bg_dark, roles['primary'], 0.22)),
        'accent_gold': hex_to_rgb(roles['accent']),
        'accent_blue': hex_to_rgb(roles['primary']),
        'white': hex_to_rgb(roles['on_dark']),
        'light_blue': hex_to_rgb(ensure_contrast(roles['secondary'], bg_dark, [roles['on_dark']], 3.0)),
        'gray': hex_to_rgb(ensure_contrast(roles['text_light'], bg_dark, ['9AA4B2'], 3.0)),
        'dark_gray': hex_to_rgb(mix(bg_dark, 'FFFFFF', 0.28)),
    }


def palette_mbe(roles):
    """generate_mbe.py 的 MBE 键位（白底 + 黑描边 + 高饱和撞色）。"""
    bg = roles['bg_light']
    border = ensure_contrast(roles['text'], bg, ['000000'], 7.0)
    slots = roles.get('slots', {})

    def accent(*keys):
        for k in keys:
            if slots.get(k):
                return ensure_contrast(slots[k], bg, [], 2.0)
        return roles['accent']

    return {
        'bg': hex_to_rgb(bg),
        'border': hex_to_rgb(border),
        'text': hex_to_rgb(roles['text']),
        'yellow': hex_to_rgb(accent('accent3', 'accent5')),
        'purple': hex_to_rgb(accent('accent4', 'accent6')),
        'red': hex_to_rgb(accent('accent2', 'accent1')),
        'blue': hex_to_rgb(roles['primary']),
        'green': hex_to_rgb(mix(roles['accent'], roles['primary'], 0.5)),
        'orange': hex_to_rgb(mix(roles['accent'], 'F59E0B', 0.5)),
        'gray': hex_to_rgb(roles['text_light']),
    }


def palette_ppt(roles, name='模板'):
    """generate_ppt.py 的 COLOR_PALETTES 单条键位。"""
    return {
        'name': name,
        'primary': hex_to_rgb(roles['primary']),
        'secondary': hex_to_rgb(roles['secondary']),
        'accent': hex_to_rgb(roles['accent']),
        'bg_light': hex_to_rgb(roles['bg_light']),
        'bg_dark': hex_to_rgb(roles['bg_dark']),
        'text': hex_to_rgb(roles['text']),
        'text_light': hex_to_rgb(roles['text_light']),
    }


# ------------------------------------------------------------------ #
# 演示文稿构造：真正的母版继承
# ------------------------------------------------------------------ #

def strip_slides(prs):
    """移除母版自带的全部幻灯片，只保留 slideMaster / slideLayout / theme。"""
    try:
        sld_id_lst = prs.slides._sldIdLst
    except AttributeError:  # pragma: no cover
        return 0

    removed = 0
    for sld_id in list(sld_id_lst):
        rid = sld_id.get(qn('r:id')) if qn else None
        if rid:
            try:
                prs.part.drop_rel(rid)
            except Exception:
                pass
        sld_id_lst.remove(sld_id)
        removed += 1
    return removed


def pick_blank_layout(prs):
    """挑选空白版式：优先无占位符者，其次名字像空白的，最后退化为占位符最少者。"""
    layouts = list(prs.slide_layouts)
    if not layouts:
        return prs.slide_layouts[0]

    for layout in layouts:
        if len(layout.placeholders) == 0:
            return layout

    blank_names = ('blank', '空白', 'empty', 'none', '唯空白')
    for layout in layouts:
        if (layout.name or '').strip().lower() in blank_names:
            return layout

    return min(layouts, key=lambda l: len(l.placeholders))


def list_layouts(prs):
    return [(i, layout.name, len(layout.placeholders)) for i, layout in enumerate(prs.slide_layouts)]


def create_prs(template_path=None, size=DEFAULT_SIZE, verbose=True):
    """
    构造 Presentation：
      - 给了 template_path -> 以模板为基底（继承母版/版式/主题），并清空其原有幻灯片
      - 否则               -> 新建 16:9 演示文稿
    """
    if template_path:
        prs = Presentation(template_path)
        removed = strip_slides(prs)

        if verbose:
            w = prs.slide_width / 914400.0
            h = prs.slide_height / 914400.0
            print('🎨 已以用户模板为基底：{}'.format(template_path))
            print('   继承：slideMaster {} 个 / slideLayout {} 个 / theme ✅'.format(
                len(prs.slide_masters), len(prs.slide_layouts)))
            print('   已清空模板自带幻灯片 {} 页，画布 {:.2f}in × {:.2f}in'.format(removed, w, h))
            if abs((w / h) - (16.0 / 9.0)) > 0.05:
                print('   ⚠️ 模板比例非 16:9，脚本内置坐标按 16:9 设计，版式可能有偏移')
        return prs

    prs = Presentation()
    prs.slide_width = Inches(size[0])
    prs.slide_height = Inches(size[1])
    return prs


def prepare(template_ref=None, keep_colors=False, verbose=True):
    """
    各脚本的统一入口：解析模板 -> 构造 prs -> 取主题角色与字体。

    返回 dict(prs, roles, font, theme_name, meta, template_path)
    """
    tpl_path = None
    meta = None
    if template_ref:
        tpl_path = resolve_template(template_ref)
        meta = next((t for t in list_templates() if t['path'] == tpl_path), None)

    prs = create_prs(tpl_path, verbose=verbose)

    roles = None
    font = None
    theme_name = None
    if tpl_path:
        theme = read_theme(tpl_path)
        theme_name = theme.get('name')
        if not keep_colors:
            roles = make_palette(theme)
        font = theme.get('fonts', {}).get('major') or theme.get('fonts', {}).get('minor')
        if verbose:
            print('   主题：{} | 字体：{}'.format(theme_name or '未命名', font or DEFAULT_FONT))
            if roles:
                print('   配色：primary #{} / secondary #{} / accent #{} / bg #{}'.format(
                    roles['primary'], roles['secondary'], roles['accent'], roles['bg_light']))
            elif keep_colors:
                print('   （--keep-colors：保留脚本内置配色，仅继承母版/版式/字体）')

    return {
        'prs': prs,
        'roles': roles,
        'font': font,
        'theme_name': theme_name,
        'meta': meta,
        'template_path': tpl_path,
    }


# ------------------------------------------------------------------ #
# 命令行
# ------------------------------------------------------------------ #

def build_arg_parser(default_output, description=''):
    import argparse
    parser = argparse.ArgumentParser(description=description)
    parser.add_argument('-m', '--template', default=None,
                        help='用户模板名 / 文件名 / 路径（见 templates/ 目录）')
    parser.add_argument('-o', '--output', default=default_output, help='输出 .pptx 路径')
    parser.add_argument('-t', '--title', '--topic', dest='title', default=None, help='标题（部分脚本使用）')
    parser.add_argument('--keep-colors', action='store_true',
                        help='继承母版/版式/字体，但保留脚本内置配色')
    parser.add_argument('--list-templates', action='store_true', help='列出 templates/ 下可用模板后退出')
    return parser


def run_frontend(default_output, run, description=''):
    """
    各脚本 __main__ 的公共前台：
      - 处理 --list-templates
      - 解析参数
      - 调用 run(opts) 并打印结果
    """
    parser = build_arg_parser(default_output, description)
    opts = parser.parse_args()

    if opts.list_templates:
        print_templates()
        return 0

    try:
        output = run(opts)
    except ValueError as err:
        print('❌ 模板加载失败：{}'.format(err), file=sys.stderr)
        print('   可运行 `python3 {} --list-templates` 查看可用模板，或去掉 --template。'.format(
            os.path.basename(sys.argv[0])), file=sys.stderr)
        return 1

    print('\n📁 文件位置: {}'.format(output))
    return 0


# ------------------------------------------------------------------ #
# 模块自测
# ------------------------------------------------------------------ #

if __name__ == '__main__':
    if '--list-templates' in sys.argv or '--list' in sys.argv:
        print_templates()
        sys.exit(0)

    if len(sys.argv) > 1 and not sys.argv[1].startswith('-'):
        info = prepare(sys.argv[1], keep_colors='--keep-colors' in sys.argv)
        print('版式列表（序号 / 名称 / 占位符数）：')
        for row in list_layouts(info['prs']):
            print('   ', row)
        sys.exit(0)

    print(__doc__)
    print('用法:')
    print('  python3 pptx_template.py --list-templates')
    print('  python3 pptx_template.py <模板名或路径>   # 检查继承结果与版式')
